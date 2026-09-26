# -*- coding: utf-8 -*-
"""自带检查脚本：验证程序能正确处理各种字段写法。

用法：在本文件夹里执行  python 自测.py
所有用例通过会打印「全部通过」。
"""

import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from wordfill import engine  # noqa: E402

TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_临时")
FAILS = []


def check(label, condition, extra=""):
    print(("  [OK] " if condition else "  [!!] ") + label + ("" if condition else "  " + str(extra)))
    if not condition:
        FAILS.append(label)


def text_of(path):
    doc = Document(path)
    return "\n".join(
        "".join(node.text or "" for node in para.iter(qn("w:t")))
        for para in doc.element.body.iter(qn("w:p"))
    )


def build_template(path):
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("甲方：")
    para.add_run("{{")
    para.add_run("甲方名称")
    para.add_run("}}")
    doc.add_paragraph("金额：{{金额|人民币大写}}（{{金额|0.00}}）")
    doc.add_paragraph("日期：{{日期|yyyy年M月d日}}")
    doc.add_paragraph("{{#if 加急}}")
    doc.add_paragraph("加急处理")
    doc.add_paragraph("{{else}}")
    doc.add_paragraph("常规处理")
    doc.add_paragraph("{{/if}}")
    doc.add_paragraph("备注：{{#if 备注}}{{备注}}{{else}}无{{/if}}")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "序号"
    table.cell(0, 1).text = "姓名"
    table.cell(0, 2).text = "金额"
    table.cell(1, 0).text = "{{#each 明细}}{{_index}}"
    table.cell(1, 1).text = "{{姓名}}"
    table.cell(1, 2).text = "{{小计|千分位}}{{/each}}"
    doc.save(path)


def main():
    os.makedirs(TMP, exist_ok=True)
    template = os.path.join(TMP, "模板.docx")
    output = os.path.join(TMP, "结果.docx")
    build_template(template)

    info = engine.scan_template(template)
    names = [field.name for field in info.fields]
    print("扫描到字段：%s" % names)
    check("能扫描出普通字段", "甲方名称" in names)
    check("能扫描出循环", any(loop.name == "明细" for loop in info.loops))
    skeleton = engine.build_skeleton(info)
    check("能生成数据骨架", isinstance(skeleton, dict) and "甲方名称" in skeleton)

    data = {
        "甲方名称": "某某公司",
        "金额": 1234.56,
        "日期": "2026-09-17",
        "加急": True,
        "备注": "",
        "明细": [{"姓名": "李四", "小计": 12000}, {"姓名": "王五", "小计": 800}],
    }
    result = engine.fill_template(template, data, output)
    text = text_of(output)
    print("---- 生成结果 ----")
    print(text)
    check("跨 run 字段替换", "甲方：某某公司" in text)
    check("人民币大写", "壹仟贰佰叁拾肆元伍角陆分" in text)
    check("数字格式", "1234.56" in text)
    check("中文日期", "2026年9月17日" in text)
    check("条件为真", "加急处理" in text and "常规处理" not in text)
    check("行内条件", "备注：无" in text)
    check("表格循环", "李四" in text and "王五" in text and "12,000" in text)
    check("没有残留标记", "{{" not in text)
    check("统计替换次数", result.filled > 0, result.filled)

    # 整张表格由数据生成
    table_template = os.path.join(TMP, "表格模板.docx")
    table_output = os.path.join(TMP, "表格结果.docx")
    doc = Document()
    doc.add_paragraph("名单如下：")
    doc.add_paragraph("{{表格:名单}}")
    doc.save(table_template)
    engine.fill_template(
        table_template,
        {"名单": {"表头": {"姓名": "3cm", "月薪": "2.5cm"},
                  "列格式": {"月薪": "千分位"},
                  "行": [{"姓名": "李四", "月薪": 18600}, {"姓名": "王五", "月薪": 24800}]}},
        table_output,
    )
    table = Document(table_output).tables[0]
    cells = [[cell.text for cell in row.cells] for row in table.rows]
    print("---- 自动生成的表格 ----")
    for row in cells:
        print("  | " + " | ".join(row))
    check("自动表格：表头 + 2 行", len(table.rows) == 3, len(table.rows))
    check("自动表格：内容正确", cells[0] == ["姓名", "月薪"] and cells[1] == ["李四", "18,600"], cells)

    # 同一列连续相同的内容纵向合并（数据库字段清单常用）
    merge_template = os.path.join(TMP, "字段清单模板.docx")
    merge_output = os.path.join(TMP, "字段清单结果.docx")
    doc = Document()
    doc.add_paragraph("{{表格:字段清单}}")
    doc.save(merge_template)
    engine.fill_template(
        merge_template,
        {
            "字段清单": {
                "表头": {"表名": "4cm", "字段": "3cm", "说明": "6cm"},
                "合并列": ["表名"],
                "行": [
                    {"表名": "icbc_merchant", "字段": "request_source", "说明": "来源"},
                    {"表名": "icbc_merchant", "字段": "cancel_reason", "说明": "作废原因"},
                ],
            }
        },
        merge_output,
    )
    merged = Document(merge_output).tables[0]
    first_tc = merged.rows[1]._tr.findall(qn("w:tc"))[0]
    second_tc = merged.rows[2]._tr.findall(qn("w:tc"))[0]
    merge_marks = [
        (tc.find(qn("w:tcPr")).find(qn("w:vMerge")).get(qn("w:val")) or "continue")
        for tc in (first_tc, second_tc)
        if tc.find(qn("w:tcPr")) is not None and tc.find(qn("w:tcPr")).find(qn("w:vMerge")) is not None
    ]
    check("相同表名自动合并单元格", merge_marks == ["restart", "continue"], merge_marks)

    # 多行文本（SQL / 代码）在 Word 里真的换行
    lines_template = os.path.join(TMP, "多行模板.docx")
    lines_output = os.path.join(TMP, "多行结果.docx")
    doc = Document()
    doc.add_paragraph("{{sql}}")
    doc.save(lines_template)
    engine.fill_template(
        lines_template,
        {"sql": "ALTER TABLE icbc_merchant\n    ADD COLUMN request_source varchar(16);"},
        lines_output,
    )
    sql_para = Document(lines_output).paragraphs[0]
    break_count = len(sql_para._p.findall(".//" + qn("w:br")))
    check("多行文本换成真正的换行", break_count == 1, break_count)
    check("多行内容完整", "ADD COLUMN request_source" in sql_para.text, sql_para.text)

    print()
    if FAILS:
        print("失败 %d 项：%s" % (len(FAILS), FAILS))
        return 1
    print("全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
